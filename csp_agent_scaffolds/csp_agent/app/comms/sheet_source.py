"""
Read-only access to the calling sheet, limited to ONE tab.

CALLING_SHEET_SOURCE=local_xlsx (default, testing): reads only the
CALLING_SHEET_TAB tab of CALLING_SHEET_LOCAL_PATH with openpyxl in
read-only mode. Nothing is downloaded and nothing is written.

CALLING_SHEET_SOURCE=google (deployment): reads only that tab through the
Sheets API with the spreadsheets.readonly scope.

Either way the result is the same: a list of dicts keyed by the header text
of the tab's header row, so the sync logic doesn't care where rows came from.
"""
import json
import logging
import os
from typing import Any

from ..config import (
    CALLING_SHEET_SOURCE, CALLING_SHEET_LOCAL_PATH, CALLING_SHEET_TAB,
    CALLING_SHEET_SPREADSHEET_ID,
)

logger = logging.getLogger(__name__)

HEADER_MARKERS = ("csp id", "csp code", "ko code")


class SheetSourceError(RuntimeError):
    pass


def _rows_to_dicts(raw_rows: list[list[Any]]) -> list[dict[str, Any]]:
    """Find the header row (the tab has a banner row above it) and zip each
    data row against it."""
    header_idx = None
    for i, row in enumerate(raw_rows[:10]):
        cells = [str(c).strip().lower() for c in row if c is not None]
        if any(c in HEADER_MARKERS for c in cells):
            header_idx = i
            break
    if header_idx is None:
        raise SheetSourceError(f"No 'CSP ID' header row found in tab '{CALLING_SHEET_TAB}'.")

    headers = [str(c).strip() if c is not None else "" for c in raw_rows[header_idx]]
    out = []
    for row in raw_rows[header_idx + 1:]:
        row = list(row) + [None] * (len(headers) - len(row))
        record = {h: row[j] for j, h in enumerate(headers) if h}
        if any(v not in (None, "") for v in record.values()):
            out.append(record)
    return out


def _read_local_xlsx(path: str) -> list[list[Any]]:
    import openpyxl

    if not os.path.exists(path):
        raise SheetSourceError(f"Calling sheet file not found: {path}")
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        if CALLING_SHEET_TAB not in wb.sheetnames:
            raise SheetSourceError(f"Tab '{CALLING_SHEET_TAB}' not found in {path}.")
        ws = wb[CALLING_SHEET_TAB]
        return [list(r) for r in ws.iter_rows(values_only=True)]
    finally:
        wb.close()


def _read_google(sheet_id: str) -> list[list[Any]]:
    from google.oauth2 import service_account
    from googleapiclient.discovery import build

    if not sheet_id:
        raise SheetSourceError("CALLING_SHEET_SPREADSHEET_ID is not set.")
    # A service account shared (Viewer) on this one spreadsheet; separate from
    # credentials.json, which is the Gmail login.
    creds_file = os.getenv("SHEETS_SERVICE_ACCOUNT_FILE", "sheets_service_account.json")
    if not os.path.exists(creds_file):
        raise SheetSourceError(f"Service-account file not found: {creds_file}")
    with open(creds_file) as f:
        if json.load(f).get("type") != "service_account":
            raise SheetSourceError("Live sheet access needs a service-account credentials file.")
    creds = service_account.Credentials.from_service_account_file(
        creds_file, scopes=["https://www.googleapis.com/auth/spreadsheets.readonly"])
    service = build("sheets", "v4", credentials=creds, cache_discovery=False)
    result = service.spreadsheets().values().get(
        spreadsheetId=sheet_id, range=f"'{CALLING_SHEET_TAB}'").execute()   # this tab only, every column
    return result.get("values", [])


def load_calling_sheet_rows() -> list[dict[str, Any]]:
    """Rows of the calling-sheet tab as dicts. Raises SheetSourceError."""
    source = CALLING_SHEET_SOURCE.strip().lower()
    if source == "google":
        raw = _read_google(CALLING_SHEET_SPREADSHEET_ID)
    elif source == "local_xlsx":
        raw = _read_local_xlsx(CALLING_SHEET_LOCAL_PATH)
    else:
        raise SheetSourceError(f"Unknown CALLING_SHEET_SOURCE '{CALLING_SHEET_SOURCE}'.")
    rows = _rows_to_dicts(raw)
    logger.info("calling_sheet_loaded source=%s tab=%s rows=%d", source, CALLING_SHEET_TAB, len(rows))
    return rows
