"""
app/vault.py
The document vault: one folder per CSP, and file names that say what each
file is, so staff can find an expired PVR without opening anything.

  <STORAGE_ROOT>/
      INDEX.xlsx                     every CSP x document, rebuilt nightly
      1A850004_KAUSHAR_JAHAN/
          1A850004_AGREEMENT_ACTIVE_2024-05-25_to_2027-05-24.pdf
          1A850004_PVR_EXPIRED_2024-03-10_to_2025-03-09.pdf   <- current copy, expired
          1A850004_IIBF_ACTIVE_2023-11-02_lifetime.pdf
          expired/      older copies: EXPIRED, or REPLACED before they expired
          unreadable/   blurred/unreadable copies, kept for audit
          rejected/     copies a reviewer rejected
      _staging/         new files land here first, then move into place

The database decides and the file name follows: desired_path() works out
where a document's file belongs from its row, and place() moves it there.
Both are safe to run any number of times. Document.storage_path is stored
relative to STORAGE_ROOT, so the whole folder can move to another server
by changing one setting.
"""
import hashlib
import logging
import re
import secrets
import shutil
from datetime import date
from pathlib import Path
from typing import Iterable, Optional

from sqlalchemy.orm import Session

from .config import STORAGE_REQUIRE_MARKER
from .config import STORAGE_ROOT as _STORAGE_ROOT_SETTING
from .models import CSP, Agreement, AgreementEvent, Document, DocumentStatus

logger = logging.getLogger(__name__)

PROJECT_ROOT = Path(__file__).resolve().parent.parent


def _resolve_root(setting: str) -> Path:
    p = Path(setting).expanduser()
    return p if p.is_absolute() else PROJECT_ROOT / p


ROOT = _resolve_root(_STORAGE_ROOT_SETTING)
MARKER = ".csp_vault"


def check_root() -> None:
    """With STORAGE_REQUIRE_MARKER on, raise unless the vault (e.g. the rack
    server mount) is really there."""
    if STORAGE_REQUIRE_MARKER and not (ROOT / MARKER).is_file():
        raise RuntimeError(f"vault_not_mounted: {ROOT / MARKER} is missing; is the rack server mounted?")

KIND = {
    "AGREEMENT": "AGREEMENT",
    "POLICE_VERIFICATION": "PVR",
    "CHARACTER_CERTIFICATE": "PVR",
    "PVR": "PVR",
    "IIBF_CERTIFICATE": "IIBF",
    "IIBF_CERTIFICATION": "IIBF",
}
EXT_FOR_MIME = {"application/pdf": ".pdf", "image/jpeg": ".jpg", "image/png": ".png"}
SUBFOLDER = {"EXPIRED": "expired", "REPLACED": "expired", "UNREADABLE": "unreadable", "REJECTED": "rejected"}
STAGING = "_staging"
REVIEW_STATUSES = {DocumentStatus.NEEDS_APPROVAL, DocumentStatus.NEEDS_REVIEW}


def _clean(s: str, limit: int) -> str:
    return re.sub(r"[^A-Za-z0-9]+", "_", (s or "").strip()).strip("_").upper()[:limit]


def folder_name(csp_code: str, csp_name: str) -> str:
    code, name = _clean(csp_code, 20) or "UNKNOWN", _clean(csp_name, 40)
    return f"{code}_{name}" if name else code


def csp_folder(csp_code: str, csp_name: str) -> Path:
    return ROOT / folder_name(csp_code, csp_name)


def kind_of(doc_type: Optional[str]) -> str:
    t = (doc_type or "").upper()
    return KIND.get(t) or _clean(t, 30) or "DOCUMENT"


def file_state(doc: Document, today: Optional[date] = None) -> str:
    """ACTIVE, EXPIRED, REVIEW, REPLACED, UNREADABLE or REJECTED."""
    today = today or date.today()
    if doc.readability == "UNREADABLE" or doc.status == DocumentStatus.UNREADABLE:
        return "UNREADABLE"
    if doc.status == DocumentStatus.REJECTED:
        return "REJECTED"
    if doc.expiry_date is not None and doc.expiry_date < today:
        return "EXPIRED"
    if not doc.is_current:
        return "REPLACED"
    if doc.status in REVIEW_STATUSES:
        return "REVIEW"
    return "ACTIVE"


def _dates_part(doc: Document, state: str) -> str:
    if state == "UNREADABLE":
        received = (doc.uploaded_at.date() if doc.uploaded_at else date.today()).isoformat()
        return f"received_{received}"
    if doc.issue_date is None:
        return "undated"
    issue = doc.issue_date.isoformat()
    if kind_of(doc.document_type) == "IIBF" and doc.expiry_date is None:
        return f"{issue}_lifetime"
    if doc.expiry_date is not None:
        return f"{issue}_to_{doc.expiry_date.isoformat()}"
    return issue


def _ext(doc: Document, current: Optional[Path]) -> str:
    if current is not None and current.suffix.lower() in (".pdf", ".jpg", ".jpeg", ".png"):
        return current.suffix.lower()
    return EXT_FOR_MIME.get(doc.mime_type or "", ".pdf")


def desired_path(doc: Document, csp: CSP, today: Optional[date] = None,
                 current: Optional[Path] = None) -> Path:
    """Where this document's file belongs (before any _2/_3 clash suffix)."""
    state = file_state(doc, today)
    folder = csp_folder(csp.current_code, csp.name)
    if not (doc.is_current and state in ("ACTIVE", "REVIEW", "EXPIRED")):
        folder = folder / SUBFOLDER.get(state, "expired")
    code = _clean(csp.current_code, 20) or "UNKNOWN"
    name = f"{code}_{kind_of(doc.document_type)}_{state}_{_dates_part(doc, state)}"
    return folder / f"{name}{_ext(doc, current)}"


# ------------------------------------------------------------------ paths
def rel(path: Path) -> str:
    """The value kept in Document.storage_path."""
    path = Path(path)
    try:
        return path.resolve().relative_to(ROOT.resolve()).as_posix()
    except ValueError:
        return str(path.resolve())


def abs_path(storage_path: Optional[str]) -> Optional[Path]:
    """Resolve a stored path. Understands relative-to-STORAGE_ROOT paths (new),
    absolute paths, and the older 'storage/documents/...' paths that were
    relative to the project folder."""
    if not storage_path:
        return None
    p = Path(storage_path)
    if p.is_absolute():
        return p
    under_root = ROOT / p
    if under_root.exists():
        return under_root
    legacy = PROJECT_ROOT / p
    if legacy.exists():
        return legacy
    return under_root


def is_inside_vault(path: Path) -> bool:
    """True for files under STORAGE_ROOT, or under the project's old storage/
    folder (files not yet moved by scripts/reorganize_vault.py)."""
    path = Path(path).resolve()
    for root in (ROOT.resolve(), (PROJECT_ROOT / "storage").resolve()):
        if path == root or root in path.parents:
            return True
    return False


def _sha(path: Path) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def _find_by_hash(folder: Path, sha256: str) -> Optional[Path]:
    """Last resort when a row's path is stale (e.g. a crash between moving the
    file and saving the row): look for the same bytes inside the CSP folder."""
    if not folder.is_dir():
        return None
    for f in folder.rglob("*"):
        if f.is_file() and _sha(f) == sha256:
            return f
    return None


def _same_slot(current: Path, want: Path) -> bool:
    if current.parent.resolve() != want.parent.resolve() or current.suffix.lower() != want.suffix.lower():
        return False
    return current.stem == want.stem or re.fullmatch(re.escape(want.stem) + r"_\d+", current.stem) is not None


def _free(want: Path) -> Path:
    if not want.exists():
        return want
    n = 2
    while True:
        cand = want.with_name(f"{want.stem}_{n}{want.suffix}")
        if not cand.exists():
            return cand
        n += 1


# --------------------------------------------------------------- writing
def stage(data: bytes, sha256: str, mime_type: Optional[str]) -> str:
    """Write new bytes into _staging/ (never a half-written file in a CSP
    folder). Returns the storage_path to put on the Document row; place()
    then moves it into the CSP folder."""
    check_root()
    d = ROOT / STAGING
    d.mkdir(parents=True, exist_ok=True)
    final = d / f"{sha256[:16]}_{secrets.token_hex(4)}{EXT_FOR_MIME.get(mime_type or '', '.pdf')}"
    tmp = final.with_suffix(final.suffix + ".part")
    tmp.write_bytes(data)
    tmp.replace(final)
    return rel(final)


def _sync_agreement_link(db: Session, doc: Document) -> None:
    if doc.agreement_id is None:
        return
    agr = db.get(Agreement, doc.agreement_id)
    if agr is not None and agr.pdf_hash == doc.sha256:
        agr.pdf_link = doc.storage_path


def place(db: Session, doc: Document, csp: Optional[CSP] = None, today: Optional[date] = None,
          dry_run: bool = False) -> Optional[tuple[str, str]]:
    """Move one document's file to where it belongs. Returns (old, new) if it
    moved (or would move, with dry_run), else None."""
    csp = csp or db.get(CSP, doc.csp_id)
    src = abs_path(doc.storage_path)
    if src is None or not src.exists():
        found = _find_by_hash(csp_folder(csp.current_code, csp.name), doc.sha256) if doc.sha256 else None
        if found is None:
            if doc.storage_path:
                logger.warning("vault_file_missing doc=%s path=%s", doc.id, doc.storage_path)
            return None
        src = found
    want = desired_path(doc, csp, today, current=src)
    if _same_slot(src, want):
        if not dry_run and doc.storage_path != rel(src):
            doc.storage_path = rel(src)
            _sync_agreement_link(db, doc)
        return None
    dest = _free(want)
    old = doc.storage_path or str(src)
    if dry_run:
        return old, rel(dest)
    dest.parent.mkdir(parents=True, exist_ok=True)
    shutil.move(str(src), str(dest))
    doc.storage_path = rel(dest)
    _sync_agreement_link(db, doc)
    db.add(AgreementEvent(csp_id=csp.id, event_type="FILE_RENAMED", source="vault", channel="SYSTEM",
                          payload={"document_id": doc.id, "from": old, "to": doc.storage_path}))
    logger.info("vault_moved doc=%s %s -> %s", doc.id, old, doc.storage_path)
    return old, doc.storage_path


def place_csp(db: Session, csp: CSP, today: Optional[date] = None, dry_run: bool = False) -> list[tuple[str, str]]:
    """Place every file of one CSP. Current copies are placed last, so an
    older copy leaves the top-level slot before the new one arrives."""
    docs = db.query(Document).filter(Document.csp_id == csp.id, Document.storage_path.isnot(None)).all()
    docs.sort(key=lambda d: (bool(d.is_current), d.id or 0))
    moves = []
    for d in docs:
        m = place(db, d, csp, today, dry_run=dry_run)
        if m:
            moves.append(m)
    return moves


def reconcile(db: Session, csps: Optional[Iterable[CSP]] = None, today: Optional[date] = None,
              dry_run: bool = False) -> dict:
    """Bring every file in line with the database: pick the current copy per
    type, refresh expiry statuses, then rename/move files. Used nightly (so
    ACTIVE becomes EXPIRED the night a document expires) and after renames of
    a CSP in the calling sheet."""
    from .compliance import REQUIRED_TYPES, refresh_category
    from .document_service import recompute_current, sync_agreement_row

    today = today or date.today()
    if csps is None:
        ids = [i for (i,) in db.query(Document.csp_id).distinct()]
        csps = db.query(CSP).filter(CSP.id.in_(ids)).all() if ids else []
    moves: list[tuple[str, str]] = []
    n = 0
    for csp in csps:
        n += 1
        if not dry_run:
            for t in REQUIRED_TYPES:
                recompute_current(db, csp, t)
            sync_agreement_row(db, csp)
            refresh_category(db, csp, today)
            db.flush()
        moves += place_csp(db, csp, today, dry_run=dry_run)
    return {"csps": n, "moved": len(moves), "moves": moves}


def orphan_files(db: Session) -> list[Path]:
    """Files inside the vault that no Document row points to (report only)."""
    known = set()
    for (p,) in db.query(Document.storage_path).filter(Document.storage_path.isnot(None)):
        a = abs_path(p)
        if a is not None:
            known.add(a.resolve())
    out = []
    if not ROOT.is_dir():
        return out
    for f in ROOT.rglob("*"):
        if not f.is_file() or f.name == "INDEX.xlsx" or f.name.startswith(".~lock"):
            continue
        if f.resolve() not in known:
            out.append(f)
    return out


# ------------------------------------------------------------------ index
def _cell(v):
    """Text that Excel would run as a formula is prefixed with a quote."""
    if isinstance(v, str) and v[:1] in ("=", "+", "-", "@", "\t", "\r"):
        return "'" + v
    return v


def write_index(db: Session, today: Optional[date] = None) -> Path:
    """INDEX.xlsx at the top of the vault: one row per CSP, with state, dates,
    days left and the file for each of the three documents."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    from .compliance import CATEGORY_NAMES, DOC_LABELS, REQUIRED_TYPES, canonical_type
    from .models import InternalUser

    today = today or date.today()
    staff = {u.id: u.name for u in db.query(InternalUser).all()}
    by_csp: dict[int, dict[str, Document]] = {}
    for d in db.query(Document).filter(Document.is_current.is_(True)):
        by_csp.setdefault(d.csp_id, {})[canonical_type(d.document_type)] = d

    wb = Workbook()
    ws = wb.active
    ws.title = "Documents"
    head = ["CSP code", "CSP name", "RM", "DC", "Category", "Folder"]
    for t in REQUIRED_TYPES:
        label = DOC_LABELS[t][0].split(" /")[0]
        head += [f"{label}: state", f"{label}: issued", f"{label}: expires", f"{label}: days left", f"{label}: file"]
    ws.append(head)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="0077B6")
    red = PatternFill("solid", fgColor="FDE2E2")
    for csp in db.query(CSP).filter(CSP.is_active_in_calling_sheet.is_(True)).order_by(CSP.current_code):
        docs = by_csp.get(csp.id, {})
        row = [csp.current_code, csp.name, staff.get(csp.rm_id, ""), staff.get(csp.dc_id, ""),
               CATEGORY_NAMES.get(csp.category or 4, ""), folder_name(csp.current_code, csp.name)]
        for t in REQUIRED_TYPES:
            d = docs.get(t)
            if d is None:
                row += ["MISSING", None, None, None, None]
                continue
            left = (d.expiry_date - today).days if d.expiry_date else None
            p = abs_path(d.storage_path)
            row += [file_state(d, today), d.issue_date, d.expiry_date if d.expiry_date else "lifetime" if kind_of(t) == "IIBF" else None,
                    left, rel(p) if p is not None else None]
        ws.append([_cell(v) for v in row])
        for i, v in enumerate(row):
            if v in ("EXPIRED", "MISSING"):
                ws.cell(row=ws.max_row, column=i + 1).fill = red
    for i, h in enumerate(head, 1):
        ws.column_dimensions[get_column_letter(i)].width = max(12, min(48, len(h) + 4))
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = ws.dimensions

    staging = ROOT / STAGING
    staging.mkdir(parents=True, exist_ok=True)
    tmp = staging / "INDEX.xlsx.part"
    wb.save(tmp)
    final = ROOT / "INDEX.xlsx"
    tmp.replace(final)
    return final


def run_nightly(db: Session, today: Optional[date] = None) -> dict:
    """The worker's vault job: refresh expiry, rename/move files, rebuild the index."""
    check_root()
    result = reconcile(db, today=today)
    db.commit()
    write_index(db, today)
    orphans = orphan_files(db)
    staging_left = [f for f in orphans if STAGING in f.parts and not f.name.endswith(".part")]
    return {"csps": result["csps"], "moved": result["moved"], "orphan_files": len(orphans),
            "stuck_in_staging": len(staging_left)}
