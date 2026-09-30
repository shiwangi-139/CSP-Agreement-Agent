"""
app/reports.py
The Excel reports, built from the same data the dashboard shows:

  CSP_Report_<date>.xlsx       Summary, All CSPs, one tab per category,
                               Expiring in 60 days
  Contacts_<date>.xlsx         All CSP contacts (CSP, RM, DC phone/email),
                               RM & DC directory
  Contact_Gaps_<date>.xlsx     Summary, CSP contact missing, RM gaps, DC gaps,
                               RM & DC missing details, Contact changes from CSPs

The worker writes all three every morning into <vault parent>/reports/ (the last
REPORTS_KEEP_DAYS days are kept); the dashboard can also download them live.
"""
import io
import logging
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Optional

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session

from . import vault
from .compliance import DOC_LABELS, REQUIRED_TYPES, SLAB_NAMES, evaluate, slab_label
from .models import CSP, ContactChangeRequest, InternalUser, OutboundMessage, OutboundStatus

logger = logging.getLogger(__name__)

REPORTS_KEEP_DAYS = 30
EXPIRING_WITHIN_DAYS = 60
HEADER = PatternFill("solid", fgColor="0092CC")        # kiosk.eko.in primary
RED = PatternFill("solid", fgColor="FDE2E2")
AMBER = PatternFill("solid", fgColor="FDF1D6")
GREEN = PatternFill("solid", fgColor="E3F4EA")
SHORT = {"AGREEMENT": "Agreement", "POLICE_VERIFICATION": "PVR", "IIBF_CERTIFICATE": "IIBF"}
GAP_TEXT = {
    "phone_missing": "Phone missing", "email_missing": "Email missing", "not_assigned": "Not assigned",
    "rm_contact_missing": "RM phone/email missing", "dc_contact_missing": "DC phone/email missing",
}


def _gap_text(codes) -> str:
    return ", ".join(GAP_TEXT.get(c, c.replace("_", " ").capitalize()) for c in (codes or []))


def _sheet(wb: Workbook, title: str, head: list[str], rows: list[list], first: bool = False,
           fills: Optional[list[dict[int, PatternFill]]] = None):
    ws = wb.active if first else wb.create_sheet()
    ws.title = title[:31]
    ws.append(head)
    for c in ws[1]:
        c.font, c.fill = Font(bold=True, color="FFFFFF"), HEADER
        c.alignment = Alignment(vertical="center", wrap_text=True)
    for i, row in enumerate(rows):
        ws.append([vault._cell(v) for v in row])
        for col, fill in (fills[i] if fills else {}).items():
            ws.cell(row=ws.max_row, column=col).fill = fill
    for i, h in enumerate(head, 1):
        longest = max([len(str(h))] + [len(str(r[i - 1])) for r in rows[:300] if i - 1 < len(r) and r[i - 1] is not None])
        ws.column_dimensions[get_column_letter(i)].width = max(10, min(50, longest + 2))
    ws.freeze_panes = "C2" if len(head) > 3 else "A2"
    if rows:
        ws.auto_filter.ref = ws.dimensions
    return ws


def _state_fill(status: str, days_left: Optional[int]) -> Optional[PatternFill]:
    if status in ("EXPIRED", "MISSING", "UNREADABLE"):
        return RED
    if days_left is not None and days_left <= EXPIRING_WITHIN_DAYS:
        return AMBER
    return GREEN if status == "VALID" else None


def _save(wb: Workbook) -> io.BytesIO:
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


# ------------------------------------------------------------ CSP report
def csp_report(db: Session, today: Optional[date] = None) -> io.BytesIO:
    today = today or date.today()
    staff = {u.id: u for u in db.query(InternalUser).all()}
    csps = db.query(CSP).filter(CSP.is_active_in_calling_sheet.is_(True)).order_by(CSP.current_code).all()

    head = ["CSP code", "CSP name", "State", "Branch", "Phone", "Email", "RM", "DC", "Slab", "Slab name",
            "Reason"]
    for t in REQUIRED_TYPES:
        head += [f"{SHORT[t]} status", f"{SHORT[t]} issued", f"{SHORT[t]} expires", f"{SHORT[t]} days left"]
    head += ["Next action", "Folder"]

    rows, fills, by_cat, expiring = [], [], {1: [], 2: [], 3: [], 4: []}, []
    missing = {t: 0 for t in REQUIRED_TYPES}
    for c in csps:
        st = evaluate(db, c, today)
        rm, dc = staff.get(c.rm_id), staff.get(c.dc_id)
        row = [c.current_code, c.name, c.state, c.branch, c.phone, c.email, rm.name if rm else None,
               dc.name if dc else None, st.category, SLAB_NAMES[st.category], st.reason]
        fill = {}
        for t in REQUIRED_TYPES:
            s = st.docs[t]
            expires = s.expiry_date or ("lifetime" if t == "IIBF_CERTIFICATE" and s.issue_date else None)
            f = _state_fill(s.status, s.days_left)
            if f is not None:
                fill[len(row) + 1] = f
            row += [s.status, s.issue_date, expires, s.days_left]
            if s.status in ("MISSING", "UNREADABLE"):
                missing[t] += 1
            if s.status == "VALID" and s.days_left is not None and s.days_left <= EXPIRING_WITHIN_DAYS:
                expiring.append([c.current_code, c.name, rm.name if rm else None, dc.name if dc else None,
                                 DOC_LABELS[t][0], s.issue_date, s.expiry_date, s.days_left])
        row += [c.next_action_at.date() if c.next_action_at else None, vault.folder_name(c.current_code, c.name)]
        rows.append(row)
        fills.append(fill)
        by_cat[st.category].append((row, fill))

    queued = db.query(OutboundMessage).filter(OutboundMessage.status == OutboundStatus.QUEUED_FOR_REVIEW).count()
    summary = [["Generated", datetime.now().strftime("%d-%m-%Y %H:%M")], ["Active CSPs (calling sheet)", len(csps)]]
    summary += [[slab_label(k), len(v)] for k, v in by_cat.items()]
    summary += [[f"{DOC_LABELS[t][0]} missing or unreadable", missing[t]] for t in REQUIRED_TYPES]
    summary += [[f"Documents expiring within {EXPIRING_WITHIN_DAYS} days", len(expiring)],
                ["Messages waiting for approval", queued]]

    wb = Workbook()
    _sheet(wb, "Summary", ["Item", "Count"], summary, first=True)
    _sheet(wb, "All CSPs", head, rows, fills=fills)
    for k, items in by_cat.items():
        _sheet(wb, f"Slab {k} {SLAB_NAMES[k]}", head, [r for r, _ in items], fills=[f for _, f in items])
    expiring.sort(key=lambda r: r[-1])
    _sheet(wb, f"Expiring {EXPIRING_WITHIN_DAYS} days",
           ["CSP code", "CSP name", "RM", "DC", "Document", "Issued", "Expires", "Days left"], expiring,
           fills=[{8: AMBER} for _ in expiring])
    return _save(wb)


# ------------------------------------------------------- contacts & gaps
def _staff_rows(db: Session, csps: list[CSP], staff: dict) -> list[list]:
    """RM / DC directory: name, role, number of CSPs, phone, email, what is missing."""
    counts: dict[int, int] = {}
    for c in csps:
        for uid in (c.rm_id, c.dc_id):
            if uid:
                counts[uid] = counts.get(uid, 0) + 1
    people = sorted((u for u in staff.values() if u.role in ("RM", "DC")), key=lambda u: (u.role, u.name or ""))
    return [[u.name, u.role, counts.get(u.id, 0), u.phone, u.email,
             ", ".join(x for x, v in (("Phone missing", u.phone), ("Email missing", u.email)) if not v)]
            for u in people]


def contacts_report(db: Session) -> io.BytesIO:
    """Contacts: every active CSP with its own, its RM's and its DC's phone
    and email, plus the RM & DC directory. Empty cells are red."""
    staff = {u.id: u for u in db.query(InternalUser).all()}
    csps = db.query(CSP).filter(CSP.is_active_in_calling_sheet.is_(True)).order_by(CSP.current_code).all()
    head = ["CSP code", "CSP name", "Phone", "Alt phone", "WhatsApp", "Email",
            "RM", "RM phone", "RM email", "DC", "DC phone", "DC email", "State", "Branch"]
    rows, fills = [], []
    for c in csps:
        rm, dc = staff.get(c.rm_id), staff.get(c.dc_id)
        row = [c.current_code, c.name, c.phone, c.alt_phone, c.whatsapp_number, c.email,
               rm.name if rm else None, rm.phone if rm else None, rm.email if rm else None,
               dc.name if dc else None, dc.phone if dc else None, dc.email if dc else None, c.state, c.branch]
        rows.append(row)
        # Highlight the contact fields that are empty (alt phone and WhatsApp are optional).
        fills.append({i + 1: RED for i in (2, 5, 6, 7, 8, 9, 10, 11) if not row[i]})
    directory = _staff_rows(db, csps, staff)
    wb = Workbook()
    _sheet(wb, "All CSP contacts", head, rows, first=True, fills=fills)
    _sheet(wb, "RM & DC directory", ["Name", "Role", "CSPs", "Phone", "Email", "Missing"], directory,
           fills=[{6: RED} if r[-1] else {} for r in directory])
    return _save(wb)


def gaps_report(db: Session) -> io.BytesIO:
    """Gaps only: what is missing, CSP by CSP, and for RMs / DCs."""
    staff = {u.id: u for u in db.query(InternalUser).all()}
    csps = db.query(CSP).filter(CSP.is_active_in_calling_sheet.is_(True)).order_by(CSP.current_code).all()
    csp_rows, rm_rows, dc_rows = [], [], []
    for c in csps:
        g = c.contact_gaps or {}
        rm, dc = staff.get(c.rm_id), staff.get(c.dc_id)
        if g.get("csp"):
            csp_rows.append([c.current_code, c.name, c.phone, c.alt_phone, c.email, _gap_text(g["csp"]),
                             rm.name if rm else None, dc.name if dc else None])
        if g.get("rm"):
            rm_rows.append([c.current_code, c.name, rm.name if rm else None, rm.phone if rm else None,
                            rm.email if rm else None, _gap_text(g["rm"])])
        if g.get("dc"):
            dc_rows.append([c.current_code, c.name, dc.name if dc else None, dc.phone if dc else None,
                            dc.email if dc else None, _gap_text(g["dc"])])
    staff_gaps = [r for r in _staff_rows(db, csps, staff) if r[-1]]
    by_id = {c.id: c for c in csps}
    changes = [[by_id[r.csp_id].current_code if r.csp_id in by_id else r.csp_id,
                by_id[r.csp_id].name if r.csp_id in by_id else None, r.field, r.old_value, r.new_value,
                r.created_at.strftime("%d-%m-%Y %H:%M") if r.created_at else None]
               for r in db.query(ContactChangeRequest).filter_by(status="PENDING").order_by(ContactChangeRequest.id)]

    summary = [["Generated", datetime.now().strftime("%d-%m-%Y %H:%M")], ["Active CSPs (calling sheet)", len(csps)],
               ["CSPs with phone or email missing", len(csp_rows)],
               ["CSPs with no RM, or RM contact missing", len(rm_rows)],
               ["CSPs with no DC, or DC contact missing", len(dc_rows)],
               ["RMs/DCs with phone or email missing", len(staff_gaps)],
               ["Contact changes sent by CSPs, not yet checked", len(changes)]]

    wb = Workbook()
    _sheet(wb, "Summary", ["Item", "Count"], summary, first=True)
    _sheet(wb, "CSP contact missing", ["CSP code", "CSP name", "Phone", "Alt phone", "Email", "Missing", "RM", "DC"],
           csp_rows, fills=[{6: RED} for _ in csp_rows])
    _sheet(wb, "RM gaps", ["CSP code", "CSP name", "RM", "RM phone", "RM email", "Missing"], rm_rows,
           fills=[{6: RED} for _ in rm_rows])
    _sheet(wb, "DC gaps", ["CSP code", "CSP name", "DC", "DC phone", "DC email", "Missing"], dc_rows,
           fills=[{6: RED} for _ in dc_rows])
    _sheet(wb, "RM & DC missing details", ["Name", "Role", "CSPs", "Phone", "Email", "Missing"], staff_gaps,
           fills=[{6: RED} for _ in staff_gaps])
    _sheet(wb, "Contact changes from CSPs", ["CSP code", "CSP name", "Field", "Old value", "New value", "Sent on"],
           changes)
    return _save(wb)


# ------------------------------------------------------------- nightly
def reports_dir() -> Path:
    return vault.ROOT.parent / "reports"


def write_daily(db: Session, today: Optional[date] = None) -> dict:
    """Write today's reports (CSP report, contacts, gaps) and drop ones older than REPORTS_KEEP_DAYS."""
    today = today or date.today()
    d = reports_dir()
    d.mkdir(parents=True, exist_ok=True)
    written = []
    for prefix, buf in (("CSP_Report", csp_report(db, today)), ("Contacts", contacts_report(db)),
                        ("Contact_Gaps", gaps_report(db))):
        final = d / f"{prefix}_{today.isoformat()}.xlsx"
        tmp = final.with_suffix(".xlsx.part")
        tmp.write_bytes(buf.getvalue())
        tmp.replace(final)
        written.append(final.name)
    cutoff = today - timedelta(days=REPORTS_KEEP_DAYS)
    removed = 0
    for f in d.glob("*_????-??-??.xlsx"):
        try:
            if date.fromisoformat(f.stem[-10:]) < cutoff:
                f.unlink()
                removed += 1
        except ValueError:
            continue
    return {"written": written, "folder": str(d), "removed_old": removed}
