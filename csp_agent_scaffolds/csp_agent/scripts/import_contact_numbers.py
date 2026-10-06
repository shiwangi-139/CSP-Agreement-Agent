"""
Add CSP contact numbers from a separate contact sheet (columns like
"CSP Code", "Mobile Number", "Home Phone", "Work Phone") to the numbers the
agent tracks. The calling sheet that seniors maintain is NOT changed: the
numbers go to their own table (csp_extra_phones).

Ignored on purpose: Email (many are wrong), and personal details such as
date of birth, age and gender (not needed).

    python -m scripts.import_contact_numbers data/contacts/csp_contacts.xlsx            # dry run + report
    python -m scripts.import_contact_numbers data/contacts/csp_contacts.xlsx --apply    # save

The report (logs/contact_import_*.csv) lists, per number:
  NEW              not known yet: will be added
  ALREADY_ON_SHEET the calling sheet already has it for this CSP: skipped
  ALREADY_ADDED    added by an earlier import: skipped
  SHARED           the same number is given for more than one CSP: skipped
  CSP_NOT_FOUND    the CSP code is not on the calling sheet: skipped
  INVALID          not a valid Indian mobile number: skipped
"""
import argparse
import csv
import re
from collections import defaultdict
from datetime import datetime
from pathlib import Path

from app.comms.sheets_sync import clean_phones
from app.db import SessionLocal
from app.models import CSP, CspExtraPhone
from app.validation import normalize_csp_code

CODE_HEADERS = ("csp code", "csp_code", "ko code", "kocode", "code", "csp id")
PHONE_HEADERS = ("mobile number", "mobile", "mobile no", "mobile no.", "phone", "home phone", "work phone",
                 "contact number", "contact no", "whatsapp", "whatsapp number", "alternate number")


def _rows(path: Path) -> list[dict]:
    # Decide by content, not by name: a CSV saved as ".xlsx" is still a CSV
    # (real .xlsx files are zip archives and start with "PK").
    with open(path, "rb") as f:
        is_xlsx = f.read(2) == b"PK"
    if is_xlsx:
        from openpyxl import load_workbook
        ws = load_workbook(path, read_only=True, data_only=True).active
        it = ws.iter_rows(values_only=True)
        head = [str(h or "").strip() for h in next(it)]
        return [dict(zip(head, r)) for r in it if any(v not in (None, "") for v in r)]
    with open(path, newline="", encoding="utf-8-sig") as f:
        return list(csv.DictReader(f))


def _pick(head: list[str], wanted: tuple) -> list[str]:
    return [h for h in head if h.strip().lower() in wanted]


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("file")
    ap.add_argument("--apply", action="store_true")
    a = ap.parse_args()
    path = Path(a.file)
    rows = _rows(path)
    if not rows:
        raise SystemExit("The sheet is empty.")
    head = list(rows[0].keys())
    code_col = (_pick(head, CODE_HEADERS) or [None])[0]
    phone_cols = _pick(head, PHONE_HEADERS)
    if not code_col or not phone_cols:
        raise SystemExit(f"Could not find the CSP code and phone columns. Headings found: {head}")
    print(f"{len(rows)} rows · code column: {code_col!r} · phone columns: {phone_cols} · email and personal details ignored")

    db = SessionLocal()
    try:
        csps = {c.current_code: c for c in db.query(CSP)}
        by_lookup = {c.lookup_code: c for c in csps.values() if c.lookup_code}
        known = defaultdict(set)                       # csp_id -> numbers already tracked
        for c in csps.values():
            for n in (c.phone, c.alt_phone, c.whatsapp_number):
                for p in clean_phones(n):
                    known[c.id].add(p)
        added_before = defaultdict(set)
        for e in db.query(CspExtraPhone):
            added_before[e.csp_id].add(e.phone)

        # first pass: every (csp, number) in the sheet
        found = []
        owners = defaultdict(set)                      # number -> CSP codes giving it (sheet + calling sheet)
        for c in csps.values():
            for p in known[c.id]:
                owners[p].add(c.current_code)
        for r in rows:
            raw_code = str(r.get(code_col) or "").strip()
            code = normalize_csp_code(raw_code)
            csp = csps.get(code) or by_lookup.get(code)
            for col in phone_cols:
                cell = r.get(col)
                if cell in (None, ""):
                    continue
                numbers = clean_phones(cell)
                if not numbers:
                    found.append((raw_code, csp, str(cell), col, "INVALID"))
                    continue
                for p in numbers:
                    found.append((raw_code, csp, p, col, None))
                    if csp is not None:
                        owners[p].add(csp.current_code)

        report, to_add, seen = [], [], set()
        counts = defaultdict(int)
        for raw_code, csp, p, col, status in found:
            if status is None:
                if csp is None:
                    status = "CSP_NOT_FOUND"
                elif (csp.id, p) in seen:
                    continue                           # same number twice for the same CSP (e.g. Mobile = Home)
                elif p in known[csp.id]:
                    status = "ALREADY_ON_SHEET"
                elif p in added_before[csp.id]:
                    status = "ALREADY_ADDED"
                elif len(owners[p]) > 1:
                    status = "SHARED"
                else:
                    status = "NEW"
                    to_add.append(CspExtraPhone(csp_id=csp.id, phone=p, source_column=col, source_file=path.name))
                if csp is not None:
                    seen.add((csp.id, p))
            counts[status] += 1
            report.append({"csp_code": raw_code, "csp_name": csp.name if csp else "", "number": p, "column": col,
                           "status": status,
                           "also_used_by": ", ".join(sorted(owners.get(p, set()) - {csp.current_code if csp else ""}))})

        out = Path("logs") / f"contact_import_{datetime.now():%Y%m%d_%H%M%S}.csv"
        out.parent.mkdir(exist_ok=True)
        with open(out, "w", newline="") as f:
            w = csv.DictWriter(f, fieldnames=["csp_code", "csp_name", "number", "column", "status", "also_used_by"])
            w.writeheader()
            w.writerows(report)
        for k in ("NEW", "ALREADY_ON_SHEET", "ALREADY_ADDED", "SHARED", "CSP_NOT_FOUND", "INVALID"):
            print(f"  {k:17} {counts[k]}")
        shared = sorted({r["number"] for r in report if r["status"] == "SHARED"})
        if shared:
            print(f"  numbers given for more than one CSP: {len(shared)} (see the report's also_used_by column)")
        print(f"report: {out}")
        if a.apply:
            db.add_all(to_add)
            db.commit()
            print(f"saved {len(to_add)} new numbers (calling sheet unchanged)")
        else:
            print(f"dry run: {len(to_add)} numbers would be added. Run again with --apply to save.")
    finally:
        db.close()


if __name__ == "__main__":
    main()
