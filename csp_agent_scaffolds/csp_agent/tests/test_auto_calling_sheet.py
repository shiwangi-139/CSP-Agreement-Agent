"""
tests/test_auto_calling_sheet.py
Unit tests for Automated Calling Sheet Detection, Field Extraction,
Excel (.xlsx) Generation, and Promotion to Master CSP Roster.
"""

import io
from datetime import datetime, timezone
from unittest.mock import MagicMock, patch
import openpyxl

from app.calling_sheet_auto import (
    extract_calling_sheet_fields,
    promote_auto_entry_to_csp,
    promote_all_auto_entries_to_csp,
    INDIAN_STATES
)
from app.reports_xlsx import generate_auto_calling_sheet_workbook
from app.models import CSP, InboundMessage, AutoCallingSheetEntry, AgreementEvent


def test_extract_calling_sheet_fields_terminal_reset():
    """Verify field extraction for external CSP terminal reset email."""
    subject = "Request to terminal reset for KO 1A851085"
    sender = "Ramesh Verma <ramesh.verma@gmail.com>"
    body = "Please reset terminal for KO 1A851085 as soon as possible."

    fields = extract_calling_sheet_fields(subject=subject, body=body, sender=sender)

    assert fields["csp_code"] == "1A851085"
    assert fields["csp_name"] == "Ramesh Verma"
    assert fields["csp_email"] == "ramesh.verma@gmail.com"
    assert fields["request_type"] == "TERMINAL_RESET"
    assert fields["terminal_status"] == "AUTO_DETECTED_FROM_EMAIL"


def test_internal_staff_forwarding_does_not_use_staff_name():
    """Verify internal staff forwarding does not attribute internal staff as CSP."""
    subject = "Request to terminal reset for KO 1A851085"
    sender = "Ganesh Kumar <ganesh.kumar@eko.co.in>"
    body = "Please reset terminal for KO 1A851085."

    fields = extract_calling_sheet_fields(subject=subject, body=body, sender=sender)

    assert fields["csp_code"] == "1A851085"
    assert fields["csp_name"] == "CSP 1A851085"
    assert fields["csp_email"] is None


def test_extract_calling_sheet_fields_phone_and_state():
    """Verify phone, branch, state, and circle extraction."""
    subject = "Request to terminal extension for KO 1A851459"
    sender = "Sunil Sharma <sunil.sharma@gmail.com>"
    body = "Kindly extend terminal for KO 1A851459. Contact mobile: 9876543210, Branch: Patna Kiosk, State: Bihar"

    fields = extract_calling_sheet_fields(subject=subject, body=body, sender=sender)

    assert fields["csp_code"] == "1A851459"
    assert fields["phone"] == "9876543210"
    assert fields["state"] == "Bihar"
    assert fields["branch"] == "Patna Kiosk"
    assert "Bihar" in fields["circle"]
    assert fields["request_type"] == "TERMINAL_EXTENSION"


def test_extract_calling_sheet_fallback_name():
    """Verify clean name fallback from email local-part when name missing."""
    subject = "Agreement renewal documents"
    sender = "ramesh.kumar.singh@gmail.com"
    body = "Sending attached agreement renewal."

    fields = extract_calling_sheet_fields(subject=subject, body=body, sender=sender)

    assert fields["csp_name"] == "Ramesh Kumar Singh"
    assert fields["csp_email"] == "ramesh.kumar.singh@gmail.com"
    assert fields["request_type"] == "AGREEMENT_DOCUMENT"


def test_auto_calling_sheet_excel_workbook_structure():
    """Verify generated Excel workbook has exact Calling Sheet New headers."""
    db = MagicMock()

    sample_entry1 = AutoCallingSheetEntry(
        id=1,
        csp_code="1A851459",
        csp_name="Patna Kiosk",
        csp_email="patna.kiosk@eko.co.in",
        phone="9876543210",
        state="Bihar",
        branch="Patna Kiosk",
        circle="Circle Bihar",
        terminal_status="AUTO_DETECTED_FROM_EMAIL",
        request_type="TERMINAL_EXTENSION",
        source_message_id="<test-msg-1@gmail.com>",
        source_subject="Request to terminal extension for KO 1A851459",
        detected_at=datetime(2026, 9, 22, 10, 0, tzinfo=timezone.utc),
        is_promoted_to_master=False
    )
    sample_entry2 = AutoCallingSheetEntry(
        id=2,
        csp_code="1A851085",
        csp_name="Ganesh Kumar",
        csp_email="ganesh.kumar@eko.co.in",
        phone=None,
        state="Pending Verification",
        branch="Main Kiosk",
        circle="General Circle",
        terminal_status="AUTO_DETECTED_FROM_EMAIL",
        request_type="TERMINAL_RESET",
        source_message_id="<test-msg-2@gmail.com>",
        source_subject="Re: Request to terminal reset for KO 1A851085",
        detected_at=datetime(2026, 9, 22, 11, 0, tzinfo=timezone.utc),
        is_promoted_to_master=True
    )

    db.query.return_value.order_by.return_value.all.return_value = [sample_entry1, sample_entry2]

    buf = generate_auto_calling_sheet_workbook(db)
    wb = openpyxl.load_workbook(buf)

    assert "Calling Sheet New" in wb.sheetnames
    ws = wb["Calling Sheet New"]

    # Verify column headers on row 4
    headers = [ws.cell(row=4, column=c).value for c in range(1, 13)]
    expected = [
        "CSP ID", "CSP Name", "CSP Mail ID", "Mobile", "State",
        "Branch Name", "Circle (LHO)", "Terminal Status",
        "Request / Document Type", "Date Detected", "Source Email Subject",
        "Promoted to Master"
    ]
    assert headers == expected

    # Verify rows 5 and 6
    assert ws.cell(row=5, column=1).value == "1A851459"
    assert ws.cell(row=5, column=2).value == "Patna Kiosk"
    assert ws.cell(row=5, column=4).value == "9876543210"
    assert ws.cell(row=5, column=12).value == "NO"

    assert ws.cell(row=6, column=1).value == "1A851085"
    assert ws.cell(row=6, column=12).value == "YES"


def test_promote_auto_entry_to_csp():
    """Verify promotion of an auto-detected entry into master CSP table."""
    db = MagicMock()

    entry = AutoCallingSheetEntry(
        id=10,
        csp_code="1A859999",
        csp_name="Sunil Sharma",
        csp_email="sunil@eko.co.in",
        phone="9811223344",
        state="Haryana",
        branch="Gurugram Kiosk",
        circle="Circle Haryana",
        terminal_status="AUTO_DETECTED_FROM_EMAIL",
        request_type="TERMINAL_RESET",
        source_subject="Terminal reset request",
        is_promoted_to_master=False,
        promoted_csp_id=None
    )

    # Query for entry returns entry; query for existing CSP returns None
    def mock_query(model):
        m = MagicMock()
        if model == AutoCallingSheetEntry:
            m.filter.return_value.first.return_value = entry
        elif model == CSP:
            m.filter.return_value.first.return_value = None
        return m

    db.query.side_effect = mock_query

    csp = promote_auto_entry_to_csp(db, 10)

    assert entry.is_promoted_to_master is True
    assert entry.promoted_at is not None
    db.add.assert_called()
    db.commit.assert_called()

