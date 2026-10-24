"""The DC master link: one link for every working CSP; locked documents,
duplicates and the Review queue listing."""
from datetime import date, datetime, timedelta

from app import auth
from app.api import portal
from app.master_link import RENEW_WITHIN_DAYS
from app.models import CSP, Document, DocumentStatus, ManualReviewQueue, ReviewStatus
from tests.test_portal_and_api import _csp, _pvr_pdf, client  # noqa: F401  (fixture)


def _link(Session):
    from app import master_link
    s = Session()
    url = master_link.create(s, "test")
    s.commit()
    s.close()
    return url.split("dc=")[1]


def _valid_doc(Session, cid, doc_type, expiry):
    s = Session()
    s.add(Document(csp_id=cid, document_type=doc_type, status=DocumentStatus.VALID, is_current=True,
                   readability="READABLE", issue_date=expiry - timedelta(days=365), expiry_date=expiry,
                   sha256=f"sha-{cid}-{doc_type}"))
    s.commit()
    s.close()


def _upload(c, dc, code, issue, data, **kw):
    return c.post("/api/portal/upload", data={"dc": dc, "code": code, "pvr_issue_date": issue.isoformat(), **kw},
                  files={"pvr_file": ("pvr.pdf", data, "application/pdf")})


def test_page_and_context_need_a_live_link(client):
    c, Session = client
    _, code, _ = _csp(Session)
    assert "Get your upload link" in c.get("/upload?dc=nope").text
    dc = _link(Session)
    page = c.get(f"/upload?dc={dc}").text
    assert "Show details" in page and "__CONTEXT_JSON__" not in page
    r = c.post("/upload", data={"action": "dc_context", "dc": dc, "code": code})
    assert r.status_code == 200 and r.json()["context"]["csp"]["code"] == code
    assert c.post("/upload", data={"action": "dc_context", "dc": dc, "code": "1A000000"}).status_code == 404
    assert c.post("/upload", data={"action": "dc_context", "dc": "nope", "code": code}).status_code == 403
    newer = _link(Session)                                            # a new link switches the old one off
    assert c.post("/upload", data={"action": "dc_context", "dc": dc, "code": code}).status_code == 403
    assert c.post("/upload", data={"action": "dc_context", "dc": newer, "code": code}).status_code == 200


def test_document_valid_for_long_is_locked_but_one_expiring_soon_is_open(client):
    c, Session = client
    cid, code, _ = _csp(Session)
    _valid_doc(Session, cid, "POLICE_VERIFICATION", date.today() + timedelta(days=200))
    _valid_doc(Session, cid, "IIBF_CERTIFICATE", date.today() + timedelta(days=RENEW_WITHIN_DAYS - 5))
    dc = _link(Session)
    docs = {d["type"]: d for d in c.post("/upload", data={"action": "dc_context", "dc": dc, "code": code})
            .json()["context"]["docs"]}
    assert docs["POLICE_VERIFICATION"]["locked"] and "valid till" in docs["POLICE_VERIFICATION"]["lock_message"]["en"]
    assert not docs["IIBF_CERTIFICATE"]["locked"] and not docs["AGREEMENT"]["locked"]
    issue = date.today() - timedelta(days=10)
    res = _upload(c, dc, code, issue, _pvr_pdf(issue)).json()["results"][0]
    assert not res["ok"] and "valid till" in res["message"]["en"]     # the server enforces the lock too


def test_dc_upload_is_stored_counted_and_listed_for_review(client):
    c, Session = client
    cid, code, _ = _csp(Session)
    dc = _link(Session)
    issue = date.today() - timedelta(days=20)
    r = _upload(c, dc, code, issue, _pvr_pdf(issue))
    assert r.status_code == 200 and r.json()["results"][0]["ok"], r.text
    s = Session()
    d = s.query(Document).filter_by(csp_id=cid).one()
    assert d.upload_channel == "DC_MASTER_LINK" and d.status == DocumentStatus.VALID
    q = s.query(ManualReviewQueue).filter_by(document_id=d.id, status=ReviewStatus.PENDING).one()
    assert q.reason.startswith("DC upload") and "already counted" in q.reason
    s.close()


def test_file_on_record_for_another_csp_is_refused(client):
    c, Session = client
    _, code_a, _ = _csp(Session)
    cid_b, code_b, _ = _csp(Session)
    dc = _link(Session)
    issue = date.today() - timedelta(days=20)
    data = _pvr_pdf(issue)
    assert _upload(c, dc, code_a, issue, data).json()["results"][0]["ok"]
    res = _upload(c, dc, code_b, issue, data).json()["results"][0]
    assert not res["ok"] and res["message"]["en"] == portal.MSG["dup_other"][1]
    s = Session()
    assert s.query(Document).filter_by(csp_id=cid_b).count() == 0
    s.close()


def test_dc_page_lists_no_csps_and_offers_no_excel(client):
    c, Session = client
    _, code, _ = _csp(Session)
    dc = _link(Session)
    page = c.get(f"/upload?dc={dc}").text
    assert code not in page and "Portal CSP" not in page and "dc_report" not in page
    r = c.post("/upload", data={"action": "dc_report", "dc": dc})
    assert r.status_code == 200 and "PK" not in r.text[:2]          # just the get-a-link page, no file


def test_admin_downloads_one_dcs_excel_in_plain_words(client, monkeypatch):
    import io
    from openpyxl import load_workbook
    from app.models import InternalUser
    monkeypatch.setattr(auth, "ADMIN_API_KEY", "k")
    hdr = {"X-API-Key": "k", "X-Requested-With": auth.CSRF_HEADER_VALUE}
    c, Session = client
    cid, code, _ = _csp(Session)
    other, other_code, _ = _csp(Session)
    s = Session()
    dc = InternalUser(name="DC Test Person", role="DC")
    s.add(dc)
    s.flush()
    s.get(CSP, cid).dc_id = dc.id
    s.commit()
    dc_id = dc.id
    s.close()
    _valid_doc(Session, cid, "POLICE_VERIFICATION", date.today() + timedelta(days=200))
    assert any(u["id"] == dc_id and u["csps"] == 1 for u in c.get("/api/hub/master-link", headers=hdr).json()["dcs"])
    r = c.get(f"/api/hub/dc-report.xlsx?dc_id={dc_id}", headers=hdr)
    assert r.status_code == 200
    ws = load_workbook(io.BytesIO(r.content)).active
    text = [[str(v) for v in row if v is not None] for row in ws.iter_rows(values_only=True)]
    assert "DC Test Person" in text[0][0]
    row = next(r for r in text if code in r)
    assert "CSP Agreement, IIBF Certificate" in row and any(v.startswith("OK till") for v in row)
    assert not any(other_code in r for r in text)                     # another DC's CSP is not in it
    assert c.get("/api/hub/dc-report.xlsx?dc_id=999999", headers=hdr).status_code == 404


def test_admin_makes_and_switches_off_the_link(client, monkeypatch):
    monkeypatch.setattr(auth, "ADMIN_API_KEY", "k")
    hdr = {"X-API-Key": "k", "X-Requested-With": auth.CSRF_HEADER_VALUE}
    c, _ = client
    made = c.post("/api/hub/master-link", headers=hdr)
    assert made.status_code == 200 and "/upload?dc=" in made.json()["url"] and made.json()["active"]
    assert c.get("/api/hub/master-link", headers=hdr).json()["active"]
    assert c.post("/api/hub/master-link/revoke", headers=hdr).json()["revoked"] >= 1
    assert not c.get("/api/hub/master-link", headers=hdr).json()["active"]


def _rename(Session, cid, name):
    s = Session()
    s.get(CSP, cid).name = name
    s.commit()
    s.close()


def test_another_csps_document_under_a_wrong_ko_code_is_refused(client):
    c, Session = client
    cid_a, code_a, _ = _csp(Session)
    cid_b, code_b, _ = _csp(Session)
    _rename(Session, cid_a, "Ramesh Yadav")
    _rename(Session, cid_b, "Suresh Chaurasia")
    dc = _link(Session)
    issue = date.today() - timedelta(days=20)
    res = _upload(c, dc, code_a, issue, _pvr_pdf(issue, holder="SURESH CHAURASIA")).json()["results"][0]
    assert not res["ok"] and "another CSP" in res["message"]["en"] and code_b in res["message"]["en"]
    s = Session()
    assert s.query(Document).filter(Document.csp_id.in_([cid_a, cid_b])).count() == 0
    s.close()
    ok = _upload(c, dc, code_a, issue, _pvr_pdf(issue, holder="RAMESH YADAV")).json()["results"][0]
    assert ok["ok"]                                                    # this CSP's own document goes in
    s = Session()
    d = s.query(Document).filter_by(csp_id=cid_a).one()
    assert "not confirmed" not in s.query(ManualReviewQueue).filter_by(document_id=d.id).one().reason
    s.close()


def test_document_without_a_readable_name_is_taken_and_flagged(client):
    c, Session = client
    cid, code, _ = _csp(Session)
    _rename(Session, cid, "Mahendra Prajapati")
    dc = _link(Session)
    issue = date.today() - timedelta(days=20)
    assert _upload(c, dc, code, issue, _pvr_pdf(issue, holder="")).json()["results"][0]["ok"]
    s = Session()
    d = s.query(Document).filter_by(csp_id=cid).one()
    assert "not confirmed" in s.query(ManualReviewQueue).filter_by(document_id=d.id).one().reason
    s.close()


def test_dc_text_summary_counts_and_lists(client, monkeypatch):
    from app.models import InternalUser
    monkeypatch.setattr(auth, "ADMIN_API_KEY", "k")
    hdr = {"X-API-Key": "k", "X-Requested-With": auth.CSRF_HEADER_VALUE}
    c, Session = client
    cid_a, code_a, _ = _csp(Session)
    cid_b, code_b, _ = _csp(Session)
    s = Session()
    dc = InternalUser(name="DC Summary Person", role="DC")
    s.add(dc)
    s.flush()
    s.get(CSP, cid_a).dc_id = s.get(CSP, cid_b).dc_id = dc.id
    s.commit()
    dc_id = dc.id
    s.close()
    _valid_doc(Session, cid_a, "POLICE_VERIFICATION", date.today() + timedelta(days=30))
    text = c.get(f"/api/hub/dc-summary?dc_id={dc_id}", headers=hdr).json()["text"]
    print(text)
    assert "Total CSPs: *2*" in text and "Need documents: *2*" in text
    assert "Renew within 60 days: *1*" in text and code_a in text and code_b in text
    assert c.get("/api/hub/dc-summary?dc_id=999999", headers=hdr).status_code == 404


def test_dc_progress_compares_the_previous_day_with_now(client, monkeypatch):
    from app.models import InternalUser, LabelSnapshot
    monkeypatch.setattr(auth, "ADMIN_API_KEY", "k")
    hdr = {"X-API-Key": "k", "X-Requested-With": auth.CSRF_HEADER_VALUE}
    c, Session = client
    cid, _, _ = _csp(Session)
    s = Session()
    dc = InternalUser(name="DC Progress Person", role="DC")
    s.add(dc)
    s.flush()
    s.get(CSP, cid).dc_id = dc.id
    s.query(LabelSnapshot).filter(LabelSnapshot.csp_id == cid).delete()
    s.add(LabelSnapshot(day=date.today() - timedelta(days=1), csp_id=cid, category=4,
                        docs={"AGREEMENT": "MISSING", "POLICE_VERIFICATION": "MISSING", "IIBF_CERTIFICATE": "MISSING"}))
    s.commit()
    s.close()
    _valid_doc(Session, cid, "POLICE_VERIFICATION", date.today() + timedelta(days=200))   # uploaded today
    d = c.get("/api/hub/dc-progress", headers=hdr).json()
    assert d["base_day"] is not None
    row = next(r for r in d["rows"] if r["dc"] == "DC Progress Person")
    assert row == {"dc": "DC Progress Person", "csps": 1, "left_now": 1, "complete_before": 0, "complete_now": 0,
                   "valid_before": 0, "valid_now": 1, "improved": 1, "worse": 0, "need_now": 1}


def test_one_link_always_shown_never_expires_and_the_shared_one_can_be_brought_back(client, monkeypatch):
    from app.models import MasterLink
    monkeypatch.setattr(auth, "ADMIN_API_KEY", "k")
    hdr = {"X-API-Key": "k", "X-Requested-With": auth.CSRF_HEADER_VALUE}
    c, Session = client
    _, code, _ = _csp(Session)
    shared = c.post("/api/hub/master-link", headers=hdr).json()["url"]
    assert c.get("/api/hub/master-link", headers=hdr).json()["url"] == shared        # shown any time
    token = shared.split("dc=")[1]
    s = Session()
    s.query(MasterLink).filter(MasterLink.revoked_at.is_(None)).update({"expires_at": datetime(2020, 1, 1)})
    s.commit()
    s.close()
    ctx = lambda t: c.post("/upload", data={"action": "dc_context", "dc": t, "code": code}).status_code
    assert ctx(token) == 200                                                          # no expiry
    newer = c.post("/api/hub/master-link", headers=hdr).json()["url"]                 # by mistake
    assert ctx(token) == 403
    back = c.post("/api/hub/master-link/restore", headers=hdr, json={"url": shared})
    assert back.status_code == 200 and back.json()["url"] == shared
    assert ctx(token) == 200 and ctx(newer.split("dc=")[1]) == 403                     # still one link
    assert c.post("/api/hub/master-link/restore", headers=hdr, json={"url": "https://x/upload?dc=made-up"}).status_code == 404


def test_dc_live_sheet_gives_one_dcs_table_only_with_the_key(client, monkeypatch):
    from app.models import InternalUser
    monkeypatch.setattr(auth, "ADMIN_API_KEY", "k")
    hdr = {"X-API-Key": "k", "X-Requested-With": auth.CSRF_HEADER_VALUE}
    c, Session = client
    cid, code, _ = _csp(Session)
    s = Session()
    dc = InternalUser(name="DC Sheet Person", role="DC")
    s.add(dc)
    s.flush()
    s.get(CSP, cid).dc_id = dc.id
    s.commit()
    s.close()
    _valid_doc(Session, cid, "POLICE_VERIFICATION", date.today() + timedelta(days=200))
    assert c.get("/api/hub/dc-sheet-key", headers=hdr).json()["key"] is None
    key = c.post("/api/hub/dc-sheet-key", headers=hdr).json()["key"]
    ask = lambda k, n="dc sheet person": c.post("/upload", data={"action": "dc_sheet", "key": k, "dc_name": n})
    assert ask("wrong").status_code == 403 and ask("").status_code == 403
    assert ask(key, "Nobody").status_code == 404
    d = ask(key).json()
    assert d["ok"] and "DC Sheet Person" in d["title"] and d["header"][0] == "KO code"
    row = next(r for r in d["rows"] if r[0] == code)
    fills = d["fills"][d["rows"].index(row)]
    assert row[4] == "CSP Agreement, IIBF Certificate" and fills[4] == "red" and fills[6] == "green"
    newer = c.post("/api/hub/dc-sheet-key", headers=hdr).json()["key"]                  # the old key stops
    assert ask(key).status_code == 403 and ask(newer).status_code == 200
