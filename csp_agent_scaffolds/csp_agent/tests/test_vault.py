"""The document vault: file names follow the database, on the test DB and a
temporary vault folder (tests/conftest.py). No real CSP files are used."""
import hashlib
import uuid
from datetime import date, datetime, timedelta

from app import vault
from app.document_service import store_extracted_document
from app.models import CSP, AgreementEvent, Document, DocumentStatus

TODAY = date(2026, 9, 25)


def _csp(db, name="Kaushar Jahan"):
    code = f"7V{uuid.uuid4().int % 1000000:06d}"
    c = CSP(name=name, current_code=code, lookup_code=code, is_active_in_calling_sheet=True)
    db.add(c)
    db.flush()
    return c


def _ex(doc_type, issue, expiry, readability="READABLE", source="PVR_DIGITAL_SIGNATURE_DATE"):
    return {"readability": readability, "document_type": doc_type,
            "start_date": issue.isoformat() if issue else None,
            "expiry_date": expiry.isoformat() if expiry else None, "date_source": source}


def _store(db, csp, doc_type, issue, expiry, **kw):
    data = f"%PDF-1.4 {uuid.uuid4().hex}".encode()
    out = store_extracted_document(db, csp, data, "x.pdf", "application/pdf", _ex(doc_type, issue, expiry, **kw),
                                   channel="TEST")
    db.flush()
    return out.document


def _name(d):
    return vault.abs_path(d.storage_path).name


def _files(csp):
    folder = vault.csp_folder(csp.current_code, csp.name)
    return sorted(p.relative_to(folder).as_posix() for p in folder.rglob("*") if p.is_file())


# ------------------------------------------------------------ pure naming
def test_names_for_each_state():
    csp = CSP(name="Kaushar Jahan", current_code="1A850004")
    agr = Document(document_type="AGREEMENT", issue_date=date(2024, 5, 25), expiry_date=date(2027, 5, 24),
                   is_current=True, readability="READABLE", status=DocumentStatus.VALID, mime_type="application/pdf")
    p = vault.desired_path(agr, csp, TODAY)
    assert p.parent.name == "1A850004_KAUSHAR_JAHAN"
    assert p.name == "1A850004_AGREEMENT_ACTIVE_2024-05-25_to_2027-05-24.pdf"

    pvr = Document(document_type="CHARACTER_CERTIFICATE", issue_date=date(2024, 3, 10), expiry_date=date(2025, 3, 9),
                   is_current=True, readability="READABLE", status=DocumentStatus.VALID, mime_type="image/jpeg")
    p = vault.desired_path(pvr, csp, TODAY)
    assert p.parent.name == "1A850004_KAUSHAR_JAHAN"          # current copy stays at the top
    assert p.name == "1A850004_PVR_EXPIRED_2024-03-10_to_2025-03-09.jpg"

    iibf = Document(document_type="IIBF_CERTIFICATE", issue_date=date(2023, 11, 2), expiry_date=None,
                    is_current=True, readability="READABLE", status=DocumentStatus.NEEDS_APPROVAL)
    assert vault.desired_path(iibf, csp, TODAY).name == "1A850004_IIBF_REVIEW_2023-11-02_lifetime.pdf"

    old = Document(document_type="POLICE_VERIFICATION", issue_date=date(2026, 1, 5), expiry_date=date(2027, 1, 4),
                   is_current=False, readability="READABLE", status=DocumentStatus.VALID)
    p = vault.desired_path(old, csp, TODAY)
    assert p.parent.name == "expired" and p.name.startswith("1A850004_PVR_REPLACED_2026-01-05")

    bad = Document(document_type="AGREEMENT", is_current=False, readability="UNREADABLE",
                   status=DocumentStatus.UNREADABLE, uploaded_at=datetime(2026, 9, 20, 10, 0))
    p = vault.desired_path(bad, csp, TODAY)
    assert p.parent.name == "unreadable" and p.name == "1A850004_AGREEMENT_UNREADABLE_received_2026-09-20.pdf"

    undated = Document(document_type="AGREEMENT", is_current=True, readability="READABLE", status=DocumentStatus.VALID)
    assert vault.desired_path(undated, csp, TODAY).name == "1A850004_AGREEMENT_ACTIVE_undated.pdf"


# ---------------------------------------------------------- with the database
def test_active_file_is_renamed_expired_the_night_it_expires(db_session):
    csp = _csp(db_session)
    d = _store(db_session, csp, "POLICE_VERIFICATION", date.today() - timedelta(days=300),
               date.today() + timedelta(days=65))
    assert "_PVR_ACTIVE_" in _name(d)
    later = date.today() + timedelta(days=66)
    vault.reconcile(db_session, [csp], today=later)
    assert "_PVR_EXPIRED_" in _name(d) and vault.abs_path(d.storage_path).parent.name != "expired"
    assert d.status == DocumentStatus.EXPIRED
    ev = db_session.query(AgreementEvent).filter_by(csp_id=csp.id, event_type="FILE_RENAMED").all()
    assert ev and ev[-1].payload["to"] == d.storage_path


def test_renewed_copy_pushes_expired_one_into_expired_folder(db_session):
    csp = _csp(db_session)
    old = _store(db_session, csp, "POLICE_VERIFICATION", date.today() - timedelta(days=400),
                 date.today() - timedelta(days=35))
    assert "_PVR_EXPIRED_" in _name(old)
    new = _store(db_session, csp, "POLICE_VERIFICATION", date.today() - timedelta(days=3),
                 date.today() + timedelta(days=362))
    assert new.is_current and not old.is_current
    assert "_PVR_ACTIVE_" in _name(new) and vault.abs_path(new.storage_path).parent.name != "expired"
    assert vault.abs_path(old.storage_path).parent.name == "expired" and "_PVR_EXPIRED_" in _name(old)


def test_replaced_copy_becomes_expired_once_its_date_passes(db_session):
    csp = _csp(db_session)
    first = _store(db_session, csp, "POLICE_VERIFICATION", date.today() - timedelta(days=100),
                   date.today() + timedelta(days=265))
    _store(db_session, csp, "POLICE_VERIFICATION", date.today() - timedelta(days=2), date.today() + timedelta(days=363))
    assert "_PVR_REPLACED_" in _name(first)
    vault.reconcile(db_session, [csp], today=date.today() + timedelta(days=300))
    assert "_PVR_EXPIRED_" in _name(first) and vault.abs_path(first.storage_path).parent.name == "expired"


def test_unreadable_copy_goes_to_unreadable_folder(db_session):
    csp = _csp(db_session)
    d = _store(db_session, csp, "AGREEMENT", None, None, readability="UNREADABLE")
    assert vault.abs_path(d.storage_path).parent.name == "unreadable" and "_AGREEMENT_UNREADABLE_" in _name(d)
    assert not d.is_current


def test_csp_name_change_moves_the_folder_and_second_run_moves_nothing(db_session):
    csp = _csp(db_session)
    a = _store(db_session, csp, "AGREEMENT", date(2025, 1, 10), date(2028, 1, 9), source="AGREEMENT_HEADING_DATE")
    i = _store(db_session, csp, "IIBF_CERTIFICATE", date(2022, 7, 8), None, source="IIBF_DATED")
    csp.name = "Kaushar Jahan Begum"
    vault.reconcile(db_session, [csp])
    assert vault.abs_path(a.storage_path).parent.name.endswith("_KAUSHAR_JAHAN_BEGUM")
    assert _files(csp) == sorted([_name(a), _name(i)])
    assert vault.reconcile(db_session, [csp])["moved"] == 0


def test_paths_are_relative_and_stale_paths_heal_by_hash(db_session, _isolated_vault):
    csp = _csp(db_session)
    d = _store(db_session, csp, "AGREEMENT", date(2025, 1, 10), date(2028, 1, 9), source="AGREEMENT_HEADING_DATE")
    assert not d.storage_path.startswith("/") and d.storage_path.split("/")[0].startswith(csp.current_code)
    real = vault.abs_path(d.storage_path)
    d.storage_path = "somewhere/that/does/not/exist.pdf"     # e.g. a crash after the move
    vault.place(db_session, d, csp)
    assert vault.abs_path(d.storage_path) == real
    assert hashlib.sha256(real.read_bytes()).hexdigest() == d.sha256


def test_review_state_and_index_file(db_session):
    csp = _csp(db_session)
    d = _store(db_session, csp, "POLICE_VERIFICATION", date.today() - timedelta(days=10),
               date.today() + timedelta(days=355), source="MODEL_VISION_GROQ")
    assert "_PVR_REVIEW_" in _name(d)
    path = vault.write_index(db_session)
    from openpyxl import load_workbook
    ws = load_workbook(path).active
    rows = [r for r in ws.iter_rows(values_only=True) if r[0] == csp.current_code]
    assert rows and "REVIEW" in rows[0]


def test_index_escapes_formula_cells():
    assert vault._cell("=HYPERLINK(1)") == "'=HYPERLINK(1)" and vault._cell("Ramesh") == "Ramesh"
