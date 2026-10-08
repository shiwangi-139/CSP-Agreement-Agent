"""Find CSPs by the state of each document ("only the agreement expired"),
the overview counts, and the Excel download (app/api/hub.py)."""
import io
import uuid
from datetime import date, timedelta

import pytest
from fastapi.testclient import TestClient
from sqlalchemy.orm import sessionmaker

from app import auth
from app.db import get_db
from app.main import app
from app.models import CSP, Document, DocumentStatus

KEY = {"X-API-Key": "k"}
T = date.today()


@pytest.fixture()
def client(test_engine, monkeypatch):
    Session = sessionmaker(bind=test_engine)

    def _db():
        s = Session()
        try:
            yield s
        finally:
            s.close()
    app.dependency_overrides[get_db] = _db
    monkeypatch.setattr(auth, "ADMIN_API_KEY", "k")
    yield TestClient(app), Session
    app.dependency_overrides.clear()


def _csp(s, docs: dict) -> str:
    code = f"7F{uuid.uuid4().int % 1000000:06d}"
    c = CSP(name="Find " + code, current_code=code, lookup_code=code, phone="9876511111", is_active_in_calling_sheet=True)
    s.add(c)
    s.flush()
    for t, (issue, expiry, status) in docs.items():
        s.add(Document(csp_id=c.id, document_type=t, sha256=uuid.uuid4().hex, status=status, issue_date=issue,
                       expiry_date=expiry, is_current=True, readability="READABLE"))
    s.flush()
    return code


VALID = DocumentStatus.VALID
EXPIRED = DocumentStatus.EXPIRED


def test_only_the_agreement_expired(client):
    c, Session = client
    s = Session()
    only_agr = _csp(s, {"AGREEMENT": (T - timedelta(days=800), T - timedelta(days=5), EXPIRED),
                        "POLICE_VERIFICATION": (T - timedelta(days=30), T + timedelta(days=300), VALID),
                        "IIBF_CERTIFICATE": (date(2020, 1, 1), None, VALID)})
    agr_and_missing = _csp(s, {"AGREEMENT": (T - timedelta(days=800), T - timedelta(days=5), EXPIRED),
                               "IIBF_CERTIFICATE": (date(2020, 1, 1), None, VALID)})
    expiring = _csp(s, {"AGREEMENT": (T - timedelta(days=700), T + timedelta(days=20), VALID),
                        "POLICE_VERIFICATION": (T - timedelta(days=30), T + timedelta(days=300), VALID),
                        "IIBF_CERTIFICATE": (date(2020, 1, 1), None, VALID)})
    s.commit()
    s.close()
    codes = lambda q: {r["code"] for r in c.get(f"/api/hub/find?{q}&size=200", headers=KEY).json()["rows"]}
    only = codes("agr=EXPIRED&pvr=VALID&iibf=VALID")
    assert only_agr in only and agr_and_missing not in only
    anyother = codes("agr=EXPIRED")
    assert {only_agr, agr_and_missing} <= anyother
    assert agr_and_missing in codes("agr=EXPIRED&pvr=MISSING")
    assert expiring in codes("agr=EXPIRING") and expiring not in codes("agr=VALID")

    m = c.get("/api/hub/doc-matrix", headers=KEY).json()["counts"]
    assert m["AGREEMENT"]["EXPIRED"] >= 2 and m["AGREEMENT"]["EXPIRING"] >= 1 and m["POLICE_VERIFICATION"]["MISSING"] >= 1

    x = c.get("/api/hub/find.xlsx?agr=EXPIRED&pvr=VALID&iibf=VALID", headers=KEY)
    assert x.status_code == 200 and "agreement-EXPIRED" in x.headers["content-disposition"]
    from openpyxl import load_workbook
    ws = load_workbook(io.BytesIO(x.content)).active
    got = [row[0] for row in ws.iter_rows(min_row=2, values_only=True)]
    assert only_agr in got and agr_and_missing not in got
    assert ws.cell(1, 8).value == "CSP Agreement: state"


def test_the_test_csp_never_shows_on_the_dashboard(client):
    from app.models import OutboundMessage, OutboundStatus
    c, Session = client
    s = Session()
    test = CSP(name="TEST CSP (not real)", current_code="9T" + f"{uuid.uuid4().int % 1000000:06d}", state="TEST",
               phone="9876599999", is_active_in_calling_sheet=False)
    real = CSP(name="Real one", current_code="7R" + f"{uuid.uuid4().int % 1000000:06d}", phone="9876588888",
               is_active_in_calling_sheet=True)
    s.add_all([test, real])
    s.flush()
    for x in (test, real):
        s.add(OutboundMessage(csp_id=x.id, channel="WHATSAPP", recipient_role="CSP", template_name="ONBOARD_ALL",
                              destination=x.phone, status=OutboundStatus.SENT, attempts=1,
                              idempotency_key=uuid.uuid4().hex, payload_json={"body": "x"}))
    s.commit()
    ids = (test.id, real.id)
    s.close()
    rows = c.get("/api/hub/messages?channel=WHATSAPP&status=SENT&size=200", headers=KEY).json()["rows"]
    shown = {r["csp_id"] for r in rows}
    assert ids[1] in shown and ids[0] not in shown
